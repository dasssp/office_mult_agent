from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app


def test_upload_csv_and_reject_formula_injection() -> None:
    client = TestClient(app)
    valid = client.post("/files/upload", files={"file": ("data.csv", "name,value\na,1\n", "text/csv")})
    assert valid.status_code == 200
    unsafe = client.post("/files/upload", files={"file": ("bad.csv", "name,value\na,=1+1\n", "text/csv")})
    assert unsafe.status_code == 422


def test_file_metadata_contains_integrity_hash_and_delete_requires_permission() -> None:
    client = TestClient(app)
    uploaded = client.post(
        "/files/upload", files={"file": ("data.csv", "name,value\na,1\n", "text/csv")}
    )
    file_id = uploaded.json()["file_id"]
    metadata = client.get(f"/files/{file_id}/metadata")
    assert metadata.json()["sha256"]
    assert metadata.json()["row_count"] == 1
    assert client.delete(f"/files/{file_id}").status_code == 403
    deleted = client.delete(
        f"/files/{file_id}", headers={"x-permission-scopes": "file:delete"}
    )
    assert deleted.status_code == 200
    assert client.get(f"/files/{file_id}/metadata").status_code == 404


def test_rejects_content_type_mismatch_and_traversal_filename() -> None:
    client = TestClient(app)
    mismatch = client.post(
        "/files/upload", files={"file": ("data.csv", "a,b\n1,2\n", "application/json")}
    )
    assert mismatch.status_code == 422
    traversal = client.post(
        "/files/upload", files={"file": ("../data.csv", "a,b\n1,2\n", "text/csv")}
    )
    assert traversal.status_code == 422


def test_analyze_uploaded_file() -> None:
    client = TestClient(app)
    uploaded = client.post("/files/upload", files={"file": ("data.json", '[{"amount": 1}, {"amount": null}]', "application/json")})
    result = client.post(f"/analysis/files/{uploaded.json()['file_id']}")
    assert result.json()["null_counts"] == {"amount": 1}
    exported = client.post(f"/analysis/files/{uploaded.json()['file_id']}/export")
    assert exported.status_code == 200
    assert exported.json()["chart_svg"].endswith(".svg")


def test_upload_xlsx() -> None:
    workbook = Workbook()
    workbook.active.append(["amount"])
    workbook.active.append([1])
    buffer = BytesIO()
    workbook.save(buffer)
    response = TestClient(app).post("/files/upload", files={"file": ("data.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 200


def test_upload_docx() -> None:
    document = Document()
    document.add_paragraph("仅供提取的文本")
    buffer = BytesIO()
    document.save(buffer)
    response = TestClient(app).post("/files/upload", files={"file": ("note.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 200
